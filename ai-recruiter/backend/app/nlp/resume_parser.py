"""
resume_parser.py — extracts raw text from an uploaded resume file.

PDF: tries PyMuPDF first (fast, good general extraction), falls back to
pdfplumber if PyMuPDF returns suspiciously little text (common with
resumes that use text boxes / unusual layouts pdfplumber handles better).

DOCX: uses python-docx to walk paragraphs and table cells.
"""
import io

import pymupdf  # PyMuPDF
import pdfplumber
from docx import Document

from app.core.exceptions import AppError


class ResumeParsingError(AppError):
    def __init__(self, message: str = "Could not extract text from this resume file"):
        super().__init__(message, "RESUME_PARSE_ERROR", 422)


def _extract_pdf_pymupdf(file_bytes: bytes) -> str:
    text_parts = []
    with pymupdf.open(stream=file_bytes, filetype="pdf") as doc:
        for page in doc:
            text_parts.append(page.get_text())
    return "\n".join(text_parts)


def _extract_pdf_pdfplumber(file_bytes: bytes) -> str:
    text_parts = []
    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text() or ""
            text_parts.append(page_text)
    return "\n".join(text_parts)


def extract_pdf_text(file_bytes: bytes) -> str:
    try:
        text = _extract_pdf_pymupdf(file_bytes)
    except Exception:
        text = ""

    # PyMuPDF sometimes returns very little on unusual layouts — fall back.
    if len(text.strip()) < 30:
        try:
            fallback_text = _extract_pdf_pdfplumber(file_bytes)
            if len(fallback_text.strip()) > len(text.strip()):
                text = fallback_text
        except Exception:
            pass

    if len(text.strip()) < 10:
        raise ResumeParsingError("The PDF appears to contain no extractable text (it may be a scanned image).")

    return text.strip()


def extract_docx_text(file_bytes: bytes) -> str:
    try:
        document = Document(io.BytesIO(file_bytes))
    except Exception as exc:
        raise ResumeParsingError("This file could not be read as a .docx document.") from exc

    parts = [p.text for p in document.paragraphs if p.text.strip()]

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text.strip())

    text = "\n".join(parts).strip()
    if len(text) < 10:
        raise ResumeParsingError("The document appears to contain no extractable text.")
    return text


def extract_resume_text(file_bytes: bytes, filename: str) -> str:
    """Dispatch to the right extractor based on file extension."""
    lower = filename.lower()
    if lower.endswith(".pdf"):
        return extract_pdf_text(file_bytes)
    if lower.endswith(".docx"):
        return extract_docx_text(file_bytes)
    if lower.endswith(".doc"):
        raise ResumeParsingError(
            "Legacy .doc files are not supported — please re-save the resume as .docx or .pdf."
        )
    raise ResumeParsingError("Unsupported file type.")
